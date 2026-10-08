"""
把 PEP 裡面 type hint相關資訊整理成 Excel 檔，方便追蹤閱讀進度。
created by Claude Sonnet5
"""


from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

FONT = "Arial"
LAST = 200  # ranges extend past the data so added rows are still counted

CORE, ADV, RUN = "必讀核心", "進階擴充", "發佈與執行期"

# (category, pep number, topic, python version (from memory), why it matters, suggested order)
rows = [
    (CORE, 483, "The Theory of Type Hints", "n/a",
     "gradual typing、subtype 與 consistent-with 的差別；理解 `Any` 行為的基礎", 5),
    (CORE, 3107, "Function Annotations", "3.0",
     "只定義函式 annotation 的語法，沒有語意", None),
    (CORE, 484, "Type Hints", "3.5",
     "整套體系的起點：typing 模組、Optional / Union、TypeVar、Any、stub 檔、cast", 1),
    (CORE, 526, "Syntax for Variable Annotations", "3.6",
     "`x: int = 0`、ClassVar；對應練習中 `self.scores` 的標註", 2),
    (CORE, 585, "Type Hinting Generics In Standard Collections", "3.9",
     "標準容器直接當 generic：`deque[float]`、`tuple[float, float]`", 3),
    (CORE, 604, "Allow writing union types as X | Y", "3.10",
     "`float | None` 取代 `Optional[float]`", 4),

    (ADV, 544, "Protocols: Structural subtyping", "3.8",
     "靜態版 duck typing；之後把 score 抽象化時會用到", None),
    (ADV, 586, "Literal Types", "3.8", "以字面值當型別", None),
    (ADV, 589, "TypedDict", "3.8", "固定 key 的 dict 型別", None),
    (ADV, 591, "Adding a final qualifier to typing", "3.8", "Final：禁止重新賦值／覆寫", None),
    (ADV, 612, "Parameter Specification Variables", "3.10", "ParamSpec：標註 decorator 與高階函式", None),
    (ADV, 613, "Explicit Type Aliases", "3.10", "TypeAlias：明確標示型別別名", None),
    (ADV, 647, "User-Defined Type Guards", "3.10", "TypeGuard：自訂型別縮窄", None),
    (ADV, 646, "Variadic Generics", "3.11", "TypeVarTuple：可變數量的型別參數", None),
    (ADV, 655, "Marking individual TypedDict items as required or potentially-missing", "3.11",
     "Required / NotRequired", None),
    (ADV, 673, "Self Type", "3.11", "Self：回傳自身型別", None),
    (ADV, 675, "Arbitrary Literal String Type", "3.11", "LiteralString", None),
    (ADV, 681, "Data Class Transforms", "3.11", "dataclass_transform", None),
    (ADV, 695, "Type Parameter Syntax", "3.12",
     "`def f[T](x: T)` 與 `type X = ...`，省去 TypeVar 樣板", None),
    (ADV, 698, "Override Decorator for Static Typing", "3.12", "@override", None),
    (ADV, 696, "Type Defaults for Type Parameters", "3.13", "type parameter 的預設值", None),
    (ADV, 705, "TypedDict: Read-only items", "3.13", "ReadOnly", None),
    (ADV, 742, "Narrowing types with TypeIs", "3.13", "TypeIs：型別縮窄", None),

    (RUN, 561, "Distributing and Packaging Type Information", "3.7",
     "套件如何發佈型別資訊：`py.typed`、stub-only 套件", None),
    (RUN, 563, "Postponed Evaluation of Annotations", "3.7",
     "`from __future__ import annotations`：annotation 字串化／延後求值", None),
    (RUN, 649, "Deferred Evaluation Of Annotations Using Descriptors", "3.14",
     "3.14 採用的延遲求值機制；需在執行期讀 annotation（pydantic、dataclass 類工具）時重要", None),
    (RUN, 749, "Implementing PEP 649", "3.14", "PEP 649 的實作細節", None),
]

wb = Workbook()

# ---------------------------------------------------------------- list sheet
ws = wb.active
ws.title = "PEP 清單"

headers = ["分類", "PEP", "主題", "Python 版本（憑記憶，待核對）", "重點／為什麼重要",
           "建議順序", "連結", "閱讀狀態", "閱讀日期", "版本已核對", "筆記"]
widths = [14, 8, 42, 18, 58, 10, 12, 12, 13, 12, 44]

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F3864")
input_fill = PatternFill("solid", fgColor="FFF9C4")

for c, (h, w) in enumerate(zip(headers, widths), start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    cell.fill = head_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
    ws.column_dimensions[cell.column_letter].width = w
ws.row_dimensions[1].height = 34

for r, (cat, num, topic, ver, why, order) in enumerate(rows, start=2):
    values = [cat, num, topic, ver, why, order if order else "—",
              f"pep-{num:04d}", "未讀", None, "否", None]
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT, size=10)
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=c in (3, 5, 11),
                                   horizontal="center" if c in (1, 2, 4, 6, 7, 8, 9, 10) else "left")
        if c >= 8:
            cell.fill = input_fill
    link = ws.cell(row=r, column=7)
    link.hyperlink = f"https://peps.python.org/pep-{num:04d}/"
    link.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    ws.cell(row=r, column=9).number_format = "yyyy-mm-dd"
    ws.row_dimensions[r].height = 32

ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:K{len(rows) + 1}"

dv_status = DataValidation(type="list", formula1='"未讀,閱讀中,已讀,待重讀"', allow_blank=True)
dv_verified = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
ws.add_data_validation(dv_status)
ws.add_data_validation(dv_verified)
dv_status.add(f"H2:H{LAST}")
dv_verified.add(f"J2:J{LAST}")

for status, color in [("已讀", "C6EFCE"), ("閱讀中", "FFEB9C"), ("待重讀", "F8CBAD")]:
    ws.conditional_formatting.add(
        f"H2:H{LAST}",
        FormulaRule(formula=[f'$H2="{status}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)),
    )

# ------------------------------------------------------------- summary sheet
sm = wb.create_sheet("進度摘要")
sm.column_dimensions["A"].width = 30
for col in "BCD":
    sm.column_dimensions[col].width = 14

L = "'PEP 清單'"

def put(cell, value, bold=False, fmt=None, fill=None, color="000000", align="left", size=10):
    c = sm[cell]
    c.value = value
    c.font = Font(name=FONT, bold=bold, size=size, color=color)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c

put("A1", "PEP 閱讀進度摘要", bold=True, size=14)

put("A3", "整體進度", bold=True, color="FFFFFF", fill="1F3864")
put("B3", "數量", bold=True, color="FFFFFF", fill="1F3864", align="center")
items = [
    ("PEP 總數", f"=COUNTA({L}!B2:B{LAST})"),
    ("未讀", f'=COUNTIF({L}!H2:H{LAST},"未讀")'),
    ("閱讀中", f'=COUNTIF({L}!H2:H{LAST},"閱讀中")'),
    ("已讀", f'=COUNTIF({L}!H2:H{LAST},"已讀")'),
    ("待重讀", f'=COUNTIF({L}!H2:H{LAST},"待重讀")'),
]
for i, (label, formula) in enumerate(items, start=4):
    put(f"A{i}", label)
    put(f"B{i}", formula, align="center", fmt="0")
put("A9", "完成率（已讀 ÷ 總數）", bold=True)
put("B9", "=IF(B4=0,0,B7/B4)", bold=True, align="center", fmt="0.0%")
put("A10", "版本欄已核對")
put("B10", f'=COUNTIF({L}!J2:J{LAST},"是")', align="center", fmt="0")

for col, h in zip("ABCD", ["分類", "總數", "已讀", "完成率"]):
    put(f"{col}12", h, bold=True, color="FFFFFF", fill="1F3864", align="center" if col != "A" else "left")
for i, cat in enumerate([CORE, ADV, RUN], start=13):
    put(f"A{i}", cat)
    put(f"B{i}", f"=COUNTIF({L}!A2:A{LAST},A{i})", align="center", fmt="0")
    put(f"C{i}", f'=COUNTIFS({L}!A2:A{LAST},A{i},{L}!H2:H{LAST},"已讀")', align="center", fmt="0")
    put(f"D{i}", f"=IF(B{i}=0,0,C{i}/B{i})", align="center", fmt="0.0%")
put("A16", "合計", bold=True)
put("B16", "=SUM(B13:B15)", bold=True, align="center", fmt="0")
put("C16", "=SUM(C13:C15)", bold=True, align="center", fmt="0")
put("D16", "=IF(B16=0,0,C16/B16)", bold=True, align="center", fmt="0.0%")

put("A18", "使用說明", bold=True, color="FFFFFF", fill="1F3864")
for col in "BCD":
    sm[f"{col}18"].fill = PatternFill("solid", fgColor="1F3864")
notes = [
    ("A19", "請填寫「PEP 清單」中黃底的欄位：閱讀狀態（H）、閱讀日期（I）、版本已核對（J）、筆記（K）；其他欄位不用動。"),
    ("A20", "填寫範例：閱讀狀態＝已讀｜閱讀日期＝2026-10-08｜版本已核對＝是｜筆記＝「Optional[X] 等同 X | None」"),
    ("A22", "注意：Python 版本欄是憑記憶列出的，引用前請以每份 PEP 開頭的 Python-Version 欄位為準，核對後在 J 欄標「是」。"),
    ("A23", "PEP 是決策紀錄，不是現行規格；現行權威是 typing.python.org 上的 Specification for the Python type system。"),
]
for cell, text in notes:
    row = int(cell[1:])
    put(cell, text)
    sm.merge_cells(f"A{row}:D{row}")
    sm.row_dimensions[row].height = 44

wb.save("pep_type_hint_reading_tracker.xlsx")
print("saved")